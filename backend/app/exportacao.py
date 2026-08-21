"""
Geração de arquivos de exportação (CSV, Excel, PDF) a partir de uma lista já
filtrada/pontuada de `schemas.LicitacaoSaida` (ver `_buscar_licitacoes_pontuadas`
em main.py). Cada função devolve os bytes prontos pra ir direto numa
StreamingResponse — quem decide o filtro e o limite de itens é o endpoint.
"""

import csv
import io
import re

from fpdf import FPDF
from fpdf.enums import XPos, YPos
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

CABECALHOS = [
    "Órgão",
    "Cidade",
    "UF",
    "Objeto",
    "Valor estimado (R$)",
    "Modalidade",
    "Encerramento da proposta",
    "Link do edital",
    "Pontuação",
    "Motivos",
]


def _linha(lic) -> list:
    return [
        lic.orgao or "",
        lic.cidade or "",
        lic.uf or "",
        lic.objeto or "",
        lic.valor_estimado or 0,
        lic.modalidade or "",
        lic.data_encerramento_proposta or "",
        lic.link_edital or "",
        lic.score,
        " | ".join(lic.motivos or []),
    ]


def gerar_csv(licitacoes: list) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow(CABECALHOS)
    for lic in licitacoes:
        linha = _linha(lic)
        linha[4] = f"{linha[4]:.2f}".replace(".", ",")  # valor com vírgula decimal, padrão BR
        writer.writerow(linha)
    # utf-8-sig (com BOM) pra acentuação abrir certo quando aberto direto no Excel
    return buffer.getvalue().encode("utf-8-sig")


# Caracteres de controle (ex: vindo de textos colados/exportados de outro
# sistema no cadastro original do PNCP) que o Excel/openpyxl recusa dentro
# de uma célula — precisam ser removidos antes de gravar.
_CARACTERES_ILEGAIS_XLSX = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _limpar_para_xlsx(valor):
    if isinstance(valor, str):
        return _CARACTERES_ILEGAIS_XLSX.sub("", valor)
    return valor


def gerar_xlsx(licitacoes: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Licitações"
    ws.append(CABECALHOS)
    for celula in ws[1]:
        celula.font = Font(bold=True)

    for lic in licitacoes:
        ws.append([_limpar_para_xlsx(v) for v in _linha(lic)])

    larguras = [28, 18, 6, 50, 18, 22, 20, 40, 10, 40]
    for indice, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(indice)].width = largura

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _seguro(texto: str) -> str:
    """fpdf2 usa fontes core (latin-1) — troca qualquer caractere fora do
    latin-1 (ex: emoji vindo de algum campo do PNCP) por '?' em vez de
    quebrar a geração do PDF inteiro."""
    if not texto:
        return ""
    return texto.encode("latin-1", errors="replace").decode("latin-1")


def _formatar_valor(valor: float) -> str:
    return f"R$ {valor or 0:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def gerar_pdf(licitacoes: list, titulo: str = "Licitações exportadas") -> bytes:
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    def linha(texto: str, altura: float) -> None:
        # multi_cell por padrão deixa o cursor na borda direita da última
        # linha escrita — sem resetar pra margem esquerda, a próxima
        # chamada fica sem espaço horizontal e o fpdf2 lança FPDFException.
        pdf.multi_cell(0, altura, _seguro(texto), new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    pdf.set_font("Helvetica", "B", 15)
    pdf.set_text_color(15, 15, 17)
    linha(titulo, 8)

    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(91, 91, 98)
    linha(f"{len(licitacoes)} licitação(ões) - gerado pelo LicitTracker", 5)
    pdf.ln(3)

    for lic in licitacoes:
        cabecalho = f"{lic.orgao or 'Órgão não informado'} - {lic.uf or '—'}"
        if lic.cidade:
            cabecalho += f" - {lic.cidade}"

        pdf.set_font("Helvetica", "B", 10.5)
        pdf.set_text_color(15, 15, 17)
        linha(cabecalho, 5.5)

        pdf.set_font("Helvetica", "", 9.5)
        pdf.set_text_color(60, 60, 66)
        linha(lic.objeto or "", 5)

        prazo = lic.data_encerramento_proposta or "sem prazo informado"
        linha_meta = f"Modalidade: {lic.modalidade or '—'}  |  Valor estimado: {_formatar_valor(lic.valor_estimado)}  |  Encerramento: {prazo}"
        pdf.set_font("Helvetica", "", 8.5)
        pdf.set_text_color(139, 139, 147)
        linha(linha_meta, 5)

        if lic.link_edital:
            pdf.set_text_color(169, 135, 42)
            linha(lic.link_edital, 5)

        pdf.ln(2)
        pdf.set_draw_color(222, 222, 226)
        y = pdf.get_y()
        pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
        pdf.ln(4)

    return bytes(pdf.output())
